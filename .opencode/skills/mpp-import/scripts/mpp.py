#!/usr/bin/env python3
"""Alat MPP: audit, extract, verify.

  python3 .opencode/skills/mpp-import/scripts/mpp.py audit   MPP.xlsx
  python3 .opencode/skills/mpp-import/scripts/mpp.py extract MPP.xlsx data/mpp.json
  python3 .opencode/skills/mpp-import/scripts/mpp.py verify  data/mpp.json

Sheet di file = unit kerja, baris = "Jenis Aktivitas Unit Kerja". Kolom dipetakan lewat
nama header (bukan posisi), karena sebagian sheet tidak punya kolom "tanggal".
"""
import json
import re
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path

import openpyxl

# Sheet di file -> id kategori di aplikasi (src/lib/categories.ts).
CATEGORIES = {
    "Tata Usaha": "tata-usaha",
    "Kesekretariatan": "kesekretariatan",
    "Keuangan": "keuangan",
    "Sarpras": "sarpras",
    "SDI": "kepegawaian",
    "Humas": "humas",
    "Laboran": "laboran",
    "IT": "it",
    "Admin Prodi": "admin-prodi",
    "Kerumahtanggaan": "kerumahtanggaan",
    "Kerjasama": "kerjasama",
}

# Periode di file -> label di aplikasi (src/lib/activities.ts).
PERIODES = {
    "harian": "Harian",
    "mingguan": "Mingguan",
    "bulanan": "Bulanan",
    "triwulan": "Triwulan",
    "semester": "Semesteran",
    "semesteran": "Semesteran",
    "tahunan": "Tahunan",
    "tahun": "Tahunan",
    "kegiatan": "Kegiatan",
}

SUMMARY_PREFIXES = ("TOTAL", "WAKTU", "JUMLAH KARYAWAN")


def find_header(ws):
    """Baris header + daftar header (posisi -> nama)."""
    for row in ws.iter_rows(min_row=1, max_row=6):
        values = [str(c.value).strip() if c.value else "" for c in row]
        if any("Jenis Aktivitas" in v for v in values):
            return row[0].row, values
    raise ValueError(f"header tidak ditemukan di sheet {ws.title}")


def column_of(headers, prefix):
    for i, header in enumerate(headers):
        if header.lower().startswith(prefix.lower()):
            return i
    return None


def cell(ws, row, index):
    if index is None:
        return None
    return ws.cell(row=row, column=index + 1).value


def marked(value):
    """Kolom klasifikasi ditandai 'v'."""
    return bool(value) and "v" in str(value).strip().lower()


def day_of(value):
    """Kolom 'tanggal' berisi free text ("25 di Tiap Bulannya"); ambil angkanya.
    Kosong = 0 (tanpa tanggal)."""
    if value is None:
        return 0
    match = re.search(r"\d+", str(value))
    if not match:
        return 0
    day = int(match.group())
    return day if 1 <= day <= 31 else 0


def extract(path):
    """Semua baris aktivitas di file, dengan nilai persis seperti di file."""
    workbook = openpyxl.load_workbook(path, data_only=True)
    rows = []
    for ws in workbook.worksheets:
        if ws.title not in CATEGORIES:
            raise ValueError(f"sheet tak dikenal: {ws.title}")
        category_id = CATEGORIES[ws.title]
        header_row, headers = find_header(ws)
        cols = {
            "name": column_of(headers, "Jenis Aktivitas"),
            "date": column_of(headers, "tanggal"),
            "utama": column_of(headers, "Utama"),
            "periodik": column_of(headers, "Periodik"),
            "bukan": column_of(headers, "Bukan Tupoksi"),
            "periode": column_of(headers, "Periode"),
            "freq": column_of(headers, "Frekuensi"),
            "minutes": column_of(headers, "Durasi"),
            "sdm": column_of(headers, "Jumlah SDM"),
        }

        for r in range(header_row + 1, ws.max_row + 1):
            first = ws.cell(row=r, column=1).value
            if first is not None and str(first).strip().startswith(SUMMARY_PREFIXES):
                break  # baris ringkasan di akhir sheet

            name = cell(ws, r, cols["name"])
            if name is None or not str(name).strip():
                continue
            periode_raw = str(cell(ws, r, cols["periode"]) or "").strip()
            if periode_raw.lower() not in PERIODES:
                raise ValueError(f"{ws.title} baris {r}: periode tak dikenal {periode_raw!r}")

            # "v" di kolom Bukan Tupoksi sering berarti kolomnya geser satu.
            flags = {
                "utama": marked(cell(ws, r, cols["utama"])),
                "periodik": marked(cell(ws, r, cols["periodik"])),
                "bukan_tupoksi": marked(cell(ws, r, cols["bukan"])),
            }
            rows.append(
                {
                    "category_id": category_id,
                    "sheet": ws.title,
                    "source_row": r,
                    "name": " ".join(str(name).split()),
                    "day": day_of(cell(ws, r, cols["date"])),
                    "raw_date": "" if cell(ws, r, cols["date"]) is None else str(cell(ws, r, cols["date"])).strip(),
                    "raw_periode": periode_raw,
                    **flags,
                    "periode": PERIODES[periode_raw.lower()],
                    "frequency": int(cell(ws, r, cols["freq"]) or 0),
                    "minutes": int(cell(ws, r, cols["minutes"]) or 0),
                    "sdm": int(cell(ws, r, cols["sdm"]) or 0),
                    "flag_geser": flags["bukan_tupoksi"] and not (flags["utama"] or flags["periodik"]),
                }
            )
    return rows


# --- audit ----------------------------------------------------------------


def app_periodes():
    """Periode yang sedang dikenali aplikasi (dibaca dari source-nya)."""
    source = Path("src/lib/activities.ts")
    if not source.exists():
        return None
    match = re.search(r"export const PERIODES = \[(.*?)\]", source.read_text(), re.S)
    if not match:
        return None
    return re.findall(r'"([^"]+)"', match.group(1))


def validate_with_app(rows):
    """Jalankan aturan aplikasi (parseActivityValues) atas baris hasil ekstraksi."""
    script = Path(__file__).with_name("validate.ts")
    if not script.exists():
        return None
    with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False) as handle:
        json.dump(rows, handle, ensure_ascii=False)
        temp = handle.name
    try:
        result = subprocess.run(
            ["bun", "run", str(script), temp],
            capture_output=True,
            text=True,
        )
        if result.returncode != 0:
            return f"(validasi aplikasi gagal: {result.stderr.strip()[:200]})"
        return result.stdout.strip()
    except FileNotFoundError:
        return "(bun tidak ada — validasi aplikasi dilewati)"
    finally:
        Path(temp).unlink(missing_ok=True)


def audit(path):
    rows = extract(path)
    print(f"== {len(rows)} baris aktivitas, {len(CATEGORIES)} unit kerja ==\n")

    print("Baris per unit:")
    for category_id, count in Counter(r["category_id"] for r in rows).most_common():
        print(f"  {category_id:18s} {count}")

    print("\nRentang nilai:")
    print(f"  nama              {min(len(r['name']) for r in rows)} - {max(len(r['name']) for r in rows)} karakter")
    print(f"  frekuensi/tahun   {min(r['frequency'] for r in rows)} - {max(r['frequency'] for r in rows)}")
    print(f"  durasi (menit)    {min(r['minutes'] for r in rows)} - {max(r['minutes'] for r in rows)}")
    print(f"  jumlah SDM        {min(r['sdm'] for r in rows)} - {max(r['sdm'] for r in rows)}")
    print(f"  tanpa klasifikasi {sum(1 for r in rows if not (r['utama'] or r['periodik'] or r['bukan_tupoksi']))} baris")

    print("\nPeriode (file -> app):")
    known = app_periodes()
    variants: dict[str, Counter] = {}
    for row in rows:
        variants.setdefault(row["periode"], Counter())[row["raw_periode"]] += 1
    for label, raws in sorted(variants.items(), key=lambda kv: -sum(kv[1].values())):
        raw_text = ", ".join(f"{raw} ({count})" for raw, count in raws.most_common())
        note = ""
        if known is not None and label not in known:
            note = "  <-- belum dikenal aplikasi: tambah ke PERIODES + engine"
        print(f"  {label:12s} {sum(raws.values()):4d}  dari: {raw_text}{note}")

    shifted = [r for r in rows if r["flag_geser"]]
    if shifted:
        print("\nAnomali: tanda 'v' di kolom 'Bukan Tupoksi' (kolomnya geser):")
        for r in shifted:
            print(f"  {r['sheet']} baris {r['source_row']}: {r['name'][:60]}")

    duplicates = [k for k, c in Counter((r["category_id"], r["name"]) for r in rows).items() if c > 1]
    if duplicates:
        print(f"\nAnomali: {len(duplicates)} nama kembar dalam satu unit (importer menggabungkannya):")
        for category_id, name in duplicates:
            print(f"  {category_id}: {name[:60]}")

    dated = [r for r in rows if r["day"]]
    if dated:
        print("\nBaris bertanggal (free text -> hari):")
        for r in dated:
            print(f"  {r['sheet']} baris {r['source_row']}: {r['raw_date']!r} -> hari {r['day']}")

    print("\nValidasi aplikasi (parseActivityValues):")
    print(validate_with_app(rows) or "(tidak dijalankan)")


# --- extract ---------------------------------------------------------------


def do_extract(source, target):
    rows = extract(source)
    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    payload = [{k: v for k, v in row.items() if k not in ("raw_date", "raw_periode", "flag_geser")} for row in rows]
    target.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"{len(payload)} baris -> {target}")


# --- verify ----------------------------------------------------------------


def do_verify(target):
    import sqlite3
    from datetime import datetime, timedelta, timezone

    rows = json.loads(Path(target).read_text(encoding="utf-8"))
    db = sqlite3.connect("data/app.db")
    db.row_factory = sqlite3.Row
    # Tanggal aktivitas memakai bulan penugasan; untuk check ini bulan berjalan (Jakarta).
    jakarta = datetime.now(timezone(timedelta(hours=7)))
    month = f"{jakarta.year:04d}-{jakarta.month:02d}"
    differences = []
    for row in rows:
        found = db.execute(
            "select * from activities where category_id=? and name=?",
            (row["category_id"], row["name"]),
        ).fetchall()
        if not found:
            differences.append(f"{row['sheet']} baris {row['source_row']}: baris tidak ada di DB")
            continue
        got = found[0]
        expected = {
            "date": f"{month}-{row['day']:02d}" if row["day"] else "",
            "utama": int(row["utama"]),
            "periodik": int(row["periodik"]),
            "bukan_tupoksi": int(row["bukan_tupoksi"]),
            "periode": row["periode"],
            "frequency": row["frequency"],
            "minutes": row["minutes"],
            "sdm": row["sdm"],
            "volume": row["frequency"] * row["minutes"] * row["sdm"],
        }
        for key, value in expected.items():
            if got[key] != value:
                differences.append(f"{row['sheet']} baris {row['source_row']}: {key} file={value!r} db={got[key]!r}")

    print(f"== verifikasi {len(rows)} baris MPP vs data/app.db ==")
    print(f"selisih: {len(differences)}")
    for line in differences[:20]:
        print(f"  {line}")
    print("\nper kategori:")
    for row in db.execute("select category_id, count(*) c from activities group by category_id order by category_id"):
        print(f"  {row['category_id']:18s} {row['c']}")
    print("\ndata pemilik:")
    for table in ("users", "assignments", "activity_logs", "occurrence_notes", "posts"):
        count = db.execute(f"select count(*) c from {table}").fetchone()["c"]
        print(f"  {table:16s} {count}")


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        raise SystemExit(1)
    command = sys.argv[1]
    if command == "audit":
        audit(sys.argv[2] if len(sys.argv) > 2 else "MPP.xlsx")
    elif command == "extract":
        do_extract(sys.argv[2] if len(sys.argv) > 2 else "MPP.xlsx", sys.argv[3] if len(sys.argv) > 3 else "data/mpp.json")
    elif command == "verify":
        do_verify(sys.argv[2] if len(sys.argv) > 2 else "data/mpp.json")
    else:
        print(__doc__)
        raise SystemExit(1)


if __name__ == "__main__":
    main()
